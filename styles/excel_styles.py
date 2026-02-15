"""
تمام استایل‌ها و توابع فرمت‌دهی اکسل
Single Responsibility: فقط مسئول ظاهر
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config.colors import EXCEL_COLORS

# ─── اشیاء استایل ───
HEADER_FONT = Font(
    name='B Nazanin', bold=True, size=11,
    color=EXCEL_COLORS['header_font']
)
HEADER_FILL = PatternFill(
    start_color=EXCEL_COLORS['header_bg'],
    end_color=EXCEL_COLORS['header_bg'],
    fill_type='solid',
)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color=EXCEL_COLORS['border']),
    right=Side(style='thin', color=EXCEL_COLORS['border']),
    top=Side(style='thin', color=EXCEL_COLORS['border']),
    bottom=Side(style='thin', color=EXCEL_COLORS['border']),
)

CELL_FONT = Font(name='B Nazanin', size=10)
CELL_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

TITLE_FONT = Font(
    name='B Nazanin', bold=True, size=14,
    color=EXCEL_COLORS['header_font']
)
TITLE_FILL = PatternFill(
    start_color=EXCEL_COLORS['title_bg'],
    end_color=EXCEL_COLORS['title_bg'],
    fill_type='solid',
)

ALT_ROW_FILL = PatternFill(
    start_color=EXCEL_COLORS['alt_row'],
    end_color=EXCEL_COLORS['alt_row'],
    fill_type='solid',
)


def style_header(ws, row: int, max_col: int):
    """اعمال استایل هدر"""
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def style_data(ws, start_row: int, end_row: int, max_col: int):
    """اعمال استایل داده‌ها (zebra stripes)"""
    for r in range(start_row, end_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = CELL_FONT
            cell.alignment = CELL_ALIGN
            cell.border = THIN_BORDER
            if r % 2 == 0:
                cell.fill = ALT_ROW_FILL


def add_title(ws, text: str, row: int, max_col: int):
    """افزودن سرفصل (مرج‌شده)"""
    ws.merge_cells(
        start_row=row, start_column=1,
        end_row=row, end_column=max_col,
    )
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 35


def auto_width(ws, max_col: int):
    """تنظیم خودکار عرض ستون‌ها"""
    for c in range(1, max_col + 1):
        mx = 0
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v:
                mx = max(mx, len(str(v)) * 1.3)
        ws.column_dimensions[get_column_letter(c)].width = max(mx + 4, 13)


def write_table(ws, headers: list, data_rows: list,
                start_row: int, title_text: str = None) -> int:
    """
    نوشتن جدول کامل با هدر، استایل و عنوان اختیاری

    Returns:
        int: شماره ردیف بعد از آخرین ردیف داده (ردیف خالی بعدی)
    """
    current = start_row
    mc = len(headers)

    if title_text:
        add_title(ws, title_text, current, mc)
        current += 1

    for ci, h in enumerate(headers, 1):
        ws.cell(row=current, column=ci, value=h)
    style_header(ws, current, mc)
    current += 1

    for row_data in data_rows:
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=current, column=ci, value=val)
        current += 1

    data_start = start_row + (2 if title_text else 1)
    if current - 1 >= data_start:
        style_data(ws, data_start, current - 1, mc)

    return current
