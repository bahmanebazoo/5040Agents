"""
نقطه ورود اصلی - هماهنگ‌کننده (Orchestrator)
═══════════════════════════════════════════════
هر تغییر در یک شیت → فقط فایل آن شیت تغییر می‌کند
هر تغییر رنگ     → فقط config/colors.py تغییر می‌کند
هر تغییر استایل  → فقط styles/excel_styles.py تغییر می‌کند
هر تغییر KPI     → فقط core/kpi_calculator.py تغییر می‌کند
"""
import sys
from openpyxl import Workbook

# ─── Core ───
from core.data_loader import load_and_process
from core.kpi_calculator import calculate_all
from core.scoring import calculate_scores

# ─── Sheets ───
from sheets import (
    sheet01_dashboard,
    sheet02_agents,
    sheet03_monthly,
    sheet04_heatmap,
    sheet05_hourly,
    sheet06_weekday,
    sheet07_timerange,
    sheet08_ranking,
    sheet09_outliers,
    sheet10_daily,
    sheet11_profiles,
    sheet12_rawdata,
)


def main(input_file: str = 'getReport-12.xlsx',
         output_file: str = 'گزارش_تحلیل_جامع_عملکرد_تحویل.xlsx'):
    """اجرای اصلی"""

    # ═══ بخش ۱: داده ═══
    df = load_and_process(input_file)

    # ═══ بخش ۲: KPI ═══
    kpi = calculate_all(df)

    # ═══ بخش ۳: امتیازدهی ═══
    scoring = calculate_scores(kpi.kpi_agent)

    # ═══ بخش ۴: ساخت اکسل ═══
    print("📝 در حال ساخت فایل اکسل...")
    wb = Workbook()

    # هر شیت مستقل ساخته می‌شود
    sheet01_dashboard.build(wb, kpi)
    sheet02_agents.build(wb, kpi)
    # sheet03_monthly.build(wb, kpi)
    sheet04_heatmap.build(wb, kpi)
    sheet05_hourly.build(wb, kpi)
    sheet06_weekday.build(wb, kpi)
    sheet07_timerange.build(wb, kpi)
    sheet08_ranking.build(wb, kpi, scoring)
    sheet09_outliers.build(wb, kpi)
    sheet10_daily.build(wb, kpi)
    sheet11_profiles.build(wb, kpi, scoring)
    sheet12_rawdata.build(wb, kpi)

    # ═══ بخش ۵: ذخیره ═══
    wb.save(output_file)

    print(f"\n{'═' * 60}")
    print(f"✅ فایل خروجی با موفقیت ذخیره شد: {output_file}")
    print(f"📊 تعداد شیت‌ها: {len(wb.sheetnames)}")
    print(f"📋 شیت‌ها:")
    for i, name in enumerate(wb.sheetnames, 1):
        print(f"   {i}. {name}")
    print(f"{'═' * 60}")


if __name__ == '__main__':
    if len(sys.argv) > 1:
        main(input_file=sys.argv[1])
    else:
        main()
