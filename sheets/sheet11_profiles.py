"""
شیت ۱۱: پروفایل نمایندگی‌ها
"""
from openpyxl.styles import Font, PatternFill, Alignment

from config.constants import MONTH_NAMES, ONTIME_THRESHOLD
from config.colors import EXCEL_COLORS
from styles.excel_styles import style_header, style_data, add_title, auto_width


def build(wb, kpi, scoring):
    """ساخت شیت پروفایل تفصیلی"""
    ws = wb.create_sheet('پروفایل نمایندگی‌ها')
    ws.sheet_view.rightToLeft = True

    df = kpi.df
    mcols = kpi.month_cols_ordered
    available_months = kpi.available_months

    add_title(
        ws,
        '👤 پروفایل تفصیلی هر نمایندگی - عملکرد ماهانه',
        1, len(mcols) + 5,
    )

    current_row = 3
    sub_headers = ['ماه', 'تعداد', 'میانگین', 'میانه', 'درصد به‌موقع']
    nc = len(sub_headers)

    for agent in scoring.index:
        agent_data = df[df['نمایندگی'] == agent]

        # هدر نمایندگی
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=len(mcols) + 5,
        )
        agent_avg = round(agent_data['اختلاف_ساعت'].mean(), 1)
        agent_ontime = round(
            (agent_data['اختلاف_ساعت'] <= ONTIME_THRESHOLD).sum()
            / len(agent_data) * 100, 1
        )
        cell = ws.cell(
            row=current_row, column=1,
            value=(
                f"📌 {agent} | تعداد: {len(agent_data)} | "
                f"میانگین: {agent_avg} ساعت | "
                f"به‌موقع: {agent_ontime}% | "
                f"امتیاز: {scoring.loc[agent, 'امتیاز_نهایی']} | "
                f"{scoring.loc[agent, 'سطح_عملکرد']}"
            ),
        )
        cell.font = Font(name='B Nazanin', bold=True, size=11, color='1F3864')
        cell.fill = PatternFill(
            start_color=EXCEL_COLORS['light_blue'],
            end_color=EXCEL_COLORS['light_blue'],
            fill_type='solid',
        )
        cell.alignment = Alignment(horizontal='center', vertical='center')
        current_row += 1

        # هدر جدول
        for ci, h in enumerate(sub_headers, 1):
            ws.cell(row=current_row, column=ci, value=h)
        style_header(ws, current_row, nc)
        current_row += 1

        data_start = current_row
        for m in available_months:
            m_data = agent_data[agent_data['ماه_شمسی'] == m]
            if len(m_data) > 0:
                ws.cell(row=current_row, column=1, value=MONTH_NAMES[m])
                ws.cell(row=current_row, column=2, value=len(m_data))
                ws.cell(row=current_row, column=3,
                        value=round(m_data['اختلاف_ساعت'].mean(), 2))
                ws.cell(row=current_row, column=4,
                        value=round(m_data['اختلاف_ساعت'].median(), 2))
                ws.cell(
                    row=current_row, column=5,
                    value=round(
                        (m_data['اختلاف_ساعت'] <= ONTIME_THRESHOLD).sum()
                        / len(m_data) * 100, 1
                    ),
                )
                current_row += 1

        if current_row > data_start:
            style_data(ws, data_start, current_row - 1, nc)

        current_row += 2

    auto_width(ws, 5)
    return ws
