"""
شیت ۱۰: روند روزانه
"""
from openpyxl.chart import BarChart, LineChart, Reference

from config.colors import CHART_SPECIFIC
from styles.excel_styles import write_table, auto_width
from charts.chart_helpers import ChartPlacer, apply_single_color


def build(wb, kpi):
    """ساخت شیت روند روزانه"""
    ws = wb.create_sheet('روند روزانه')
    ws.sheet_view.rightToLeft = True

    headers = ['تاریخ', 'تعداد مرسوله', 'میانگین (ساعت)']
    rows = []
    for _, rd in kpi.daily_trend.sort_values('تاریخ').iterrows():
        rows.append([rd['تاریخ'], int(rd['تعداد']), rd['میانگین_ساعت']])

    lr = write_table(ws, headers, rows, 1, '📈 روند روزانه تحویل مرسولات')
    auto_width(ws, len(headers))

    if len(rows) > 1:
        placer = ChartPlacer(start_row=lr + 1)

        # خطی
        lc = LineChart()
        lc.title = "روند روزانه میانگین زمان تحویل"
        lc.y_axis.title = "ساعت"
        lc.style = 10
        lc.width = 35
        lc.height = 15
        lc.add_data(
            Reference(ws, min_col=3, min_row=2, max_row=lr - 1),
            titles_from_data=True,
        )
        lc.set_categories(Reference(ws, min_col=1, min_row=3, max_row=lr - 1))
        ws.add_chart(lc, placer.next_anchor())

        # ستونی تعداد
        bar = BarChart()
        bar.title = "تعداد مرسولات روزانه"
        bar.style = 10
        bar.width = 35
        bar.height = 15
        bar.add_data(
            Reference(ws, min_col=2, min_row=2, max_row=lr - 1),
            titles_from_data=True,
        )
        bar.set_categories(Reference(ws, min_col=1, min_row=3, max_row=lr - 1))
        apply_single_color(bar, CHART_SPECIFIC['daily_bar'])
        ws.add_chart(bar, placer.next_anchor())

    return ws
