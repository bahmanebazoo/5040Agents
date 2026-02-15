"""
شیت ۷: توزیع بازه‌های تحویل
"""
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule

from config.constants import TIME_RANGE_ORDER
from config.colors import CF_COLORS
from styles.excel_styles import (
    write_table, auto_width, add_title, style_header, style_data,
)
from charts.chart_helpers import ChartPlacer, colorize_category_bars, add_data_labels


def build(wb, kpi):
    """ساخت شیت توزیع بازه تحویل"""
    ws = wb.create_sheet('توزیع بازه تحویل')
    ws.sheet_view.rightToLeft = True

    df = kpi.df
    trd = kpi.time_range_dist

    # ─── جدول کلی ───
    headers = ['بازه زمانی', 'تعداد', 'درصد']
    rows = []
    total = len(df)
    for tr in TIME_RANGE_ORDER:
        if tr in trd.index:
            c = int(trd[tr])
            rows.append([tr, c, round(c / total * 100, 1)])

    lr = write_table(ws, headers, rows, 1, '📊 توزیع بازه‌های زمانی تحویل')
    auto_width(ws, len(headers))

    placer = ChartPlacer(start_row=lr + 1)

    # نمودار ستونی (چندرنگ - دسته‌بندی مفهومی = بازه)
    ch = BarChart()
    ch.type = "col"
    ch.title = "توزیع بازه‌های زمانی تحویل"
    ch.y_axis.title = "تعداد"
    ch.style = 10
    ch.width = 28
    ch.height = 15
    ch.add_data(
        Reference(ws, min_col=2, min_row=2, max_row=lr - 1),
        titles_from_data=True,
    )
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=lr - 1))
    colorize_category_bars(ch, len(rows))
    add_data_labels(ch, show_val=True)
    ws.add_chart(ch, placer.next_anchor())

    # ─── ماتریس بازه × نمایندگی ───
    sr7b = placer.current_row + 2
    ptr = kpi.pivot_timerange
    tr_cols = [c for c in TIME_RANGE_ORDER if c in ptr.columns]
    nc = len(tr_cols) + 1

    add_title(ws, '📦 ماتریس بازه تحویل × نمایندگی', sr7b, nc)
    ws.cell(row=sr7b + 1, column=1, value='نمایندگی')
    for ci, tc in enumerate(tr_cols, 2):
        ws.cell(row=sr7b + 1, column=ci, value=tc)
    style_header(ws, sr7b + 1, nc)

    for i, (agent, rd) in enumerate(ptr.iterrows()):
        r = sr7b + 2 + i
        ws.cell(row=r, column=1, value=agent)
        for ci, tc in enumerate(tr_cols, 2):
            v = rd.get(tc, 0)
            ws.cell(row=r, column=ci, value=int(v))

    lr7b = sr7b + 1 + len(ptr)
    style_data(ws, sr7b + 2, lr7b, nc)
    auto_width(ws, nc)

    ws.conditional_formatting.add(
        f'B{sr7b + 2}:{get_column_letter(nc)}{lr7b}',
        DataBarRule(
            start_type='min', end_type='max',
            color=CF_COLORS['databar_red'],
        ),
    )

    return ws
