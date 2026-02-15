"""
شیت ۴: ماتریس Heatmap (نمایندگی × ماه)
"""
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule

from config.colors import CF_COLORS
from styles.excel_styles import add_title, style_header, style_data, auto_width


def _write_matrix(ws, title_text, start_row, pivot_df, mcols, cf_type='hot'):
    """
    نوشتن یک ماتریس + Conditional Formatting

    cf_type:
        'hot'    → سبز→زرد→قرمز (کم=خوب)
        'cold'   → قرمز→زرد→سبز (کم=بد)
        'databar'→ DataBar آبی
    """
    nc = len(mcols) + 1
    add_title(ws, title_text, start_row, nc)

    hr = start_row + 1
    ws.cell(row=hr, column=1, value='نمایندگی')
    for ci, mn in enumerate(mcols, 2):
        ws.cell(row=hr, column=ci, value=mn)
    style_header(ws, hr, nc)

    for i, (agent, rd) in enumerate(pivot_df.iterrows()):
        r = hr + 1 + i
        ws.cell(row=r, column=1, value=agent)
        for ci, mn in enumerate(mcols, 2):
            v = rd.get(mn, None)
            if cf_type == 'databar':
                ws.cell(row=r, column=ci, value=int(v) if pd.notna(v) else 0)
            else:
                ws.cell(row=r, column=ci, value=v if pd.notna(v) else '-')

    lr = hr + len(pivot_df)
    style_data(ws, hr + 1, lr, nc)

    rng = f'B{hr + 1}:{get_column_letter(nc)}{lr}'
    if cf_type == 'hot':
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type='min', start_color=CF_COLORS['good_start'],
            mid_type='percentile', mid_value=50, mid_color=CF_COLORS['mid'],
            end_type='max', end_color=CF_COLORS['bad_end'],
        ))
    elif cf_type == 'cold':
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type='min', start_color=CF_COLORS['bad_end'],
            mid_type='percentile', mid_value=50, mid_color=CF_COLORS['mid'],
            end_type='max', end_color=CF_COLORS['good_start'],
        ))
    elif cf_type == 'databar':
        ws.conditional_formatting.add(rng, DataBarRule(
            start_type='min', end_type='max',
            color=CF_COLORS['databar_blue'],
        ))

    return lr + 3  # ردیف شروع ماتریس بعدی


def build(wb, kpi):
    """ساخت شیت ماتریس Heatmap"""
    ws = wb.create_sheet('ماتریس نمایندگی-ماه')
    ws.sheet_view.rightToLeft = True

    mcols = [c for c in kpi.month_cols_ordered if c in kpi.pivot_avg.columns]

    sr = 1
    sr = _write_matrix(
        ws, '🔥 Heatmap میانگین زمان تحویل (ساعت) - نمایندگی × ماه',
        sr, kpi.pivot_avg, mcols, cf_type='hot',
    )
    sr = _write_matrix(
        ws, '📦 ماتریس تعداد مرسوله - نمایندگی × ماه',
        sr, kpi.pivot_count, mcols, cf_type='databar',
    )
    sr = _write_matrix(
        ws, '✅ ماتریس درصد تحویل به‌موقع (%) - نمایندگی × ماه',
        sr, kpi.pivot_ontime, mcols, cf_type='cold',
    )
    sr = _write_matrix(
        ws, '📊 ماتریس میانه زمان تحویل (ساعت) - نمایندگی × ماه',
        sr, kpi.pivot_median, mcols, cf_type='hot',
    )

    auto_width(ws, len(mcols) + 1)
    return ws
