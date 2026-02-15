"""
شیت ۱۲: داده‌های خام
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

from styles.excel_styles import write_table, auto_width


def build(wb, kpi):
    """ساخت شیت داده‌های خام"""
    ws = wb.create_sheet('داده‌های خام')
    ws.sheet_view.rightToLeft = True

    df = kpi.df

    headers = [
        'سریال فاکتور', 'نمایندگی', 'تاریخ ورود', 'تاریخ تحویل',
        'اختلاف (ساعت)', 'وضعیت', 'بازه تحویل', 'ماه', 'روز هفته',
        'شیفت ورود', 'ساعت ورود', 'ساعت تحویل',
    ]
    rows = []
    for _, rd in df.iterrows():
        rows.append([
            rd['سریال فاکتور'], rd['نمایندگی'],
            str(rd['تاریخ ورود به نمایندگی']), str(rd['تاریخ تحویل']),
            round(rd['اختلاف_ساعت'], 2) if pd.notna(rd['اختلاف_ساعت']) else '',
            rd['وضعیت_تحویل'], rd['بازه_تحویل'],
            rd.get('نام_ماه', ''), rd.get('نام_روز_هفته', ''),
            rd.get('شیفت_ورود', ''), rd.get('ساعت_ورود', ''),
            rd.get('ساعت_تحویل', ''),
        ])

    lr = write_table(ws, headers, rows, 1, '📄 داده‌های خام با محاسبات')
    auto_width(ws, len(headers))

    # Conditional Formatting
    ws.conditional_formatting.add(
        f'E3:E{lr - 1}',
        CellIsRule(
            operator='lessThanOrEqual', formula=['24'],
            fill=PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'),
            font=Font(color='006100'),
        ),
    )
    ws.conditional_formatting.add(
        f'E3:E{lr - 1}',
        CellIsRule(
            operator='greaterThan', formula=['24'],
            fill=PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid'),
            font=Font(color='9C0006'),
        ),
    )

    # AutoFilter
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{lr - 1}"

    return ws
