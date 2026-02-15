"""
شیت ۱: داشبورد کلی
"""
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config.colors import (
    EXCEL_COLORS, PIE_STATUS_COLORS,
    PIE_TIMERANGE_COLORS,
)
from styles.excel_styles import add_title, THIN_BORDER, HEADER_FONT, HEADER_FILL, HEADER_ALIGN
from charts.chart_helpers import (
    ChartPlacer, colorize_pie, add_data_labels,
)


# ──────────────────────────────────────────────
# توابع کمکی داخلی
# ──────────────────────────────────────────────

def _style_mini_table(ws, header_row, data_count, col_start=1, col_end=2):
    """
    استایل‌دهی زیبا به جداول کوچک داده‌ای (وضعیت / بازه / شیفت)
    - هدر: پس‌زمینه آبی تیره + فونت سفید
    - ردیف‌ها: zebra stripes + حاشیه نازک
    """
    mini_border = Border(
        left=Side(style='thin', color=EXCEL_COLORS['border']),
        right=Side(style='thin', color=EXCEL_COLORS['border']),
        top=Side(style='thin', color=EXCEL_COLORS['border']),
        bottom=Side(style='thin', color=EXCEL_COLORS['border']),
    )
    # ── هدر ──
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(name='B Nazanin', bold=True, size=10, color='FFFFFF')
        cell.fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = mini_border

    # ── ردیف‌های داده ──
    alt_fill = PatternFill(
        start_color=EXCEL_COLORS['alt_row'],
        end_color=EXCEL_COLORS['alt_row'],
        fill_type='solid',
    )
    for i in range(data_count):
        r = header_row + 1 + i
        for c in range(col_start, col_end + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name='B Nazanin', size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = mini_border
            if i % 2 == 1:
                cell.fill = alt_fill


def _push_pie_plot_down(chart, offset_y=0.08):
    """
    فاصله تایتل تا بدنه نمودار دایره‌ای را افزایش می‌دهد.
    offset_y ≈ 0.10 → تقریباً ۴۸ پیکسل در نمودار ۱۲-۱۳ ردیفی
    """
    chart.layout = Layout(
        manualLayout=ManualLayout(
            x=0.02,
            y=offset_y,
            w=0.94,
            h=1.0 - offset_y - 0.03,
        )
    )


# ──────────────────────────────────────────────
# تابع اصلی ساخت شیت
# ──────────────────────────────────────────────

def build(wb, kpi):
    """ساخت شیت داشبورد کلی"""
    df = kpi.df

    ws = wb.active
    ws.title = 'داشبورد کلی'
    ws.sheet_view.rightToLeft = True

    add_title(ws, '📊 داشبورد جامع تحلیل عملکرد تحویل مرسولات', 1, 10)

    # ══════════════════════════════════════════
    # بخش ۱: KPIهای اصلی  (ردیف‌های ۳–۱۶)
    # ══════════════════════════════════════════
    kpi_main = [
        ('📦 تعداد کل مرسولات', len(df)),
        ('🏢 تعداد نمایندگی‌ها', df['نمایندگی'].nunique()),
        ('⏱️ میانگین زمان تحویل (ساعت)', round(df['اختلاف_ساعت'].mean(), 2)),
        ('⏱️ میانه زمان تحویل (ساعت)', round(df['اختلاف_ساعت'].median(), 2)),
        ('📐 انحراف معیار (ساعت)', round(df['اختلاف_ساعت'].std(), 2)),
        ('✅ درصد تحویل به‌موقع (≤۲۴ ساعت)',
         f"{round((df['اختلاف_ساعت'] <= 24).sum() / len(df) * 100, 1)}%"),
        ('❌ تعداد تأخیردار (>۲۴ ساعت)', int((df['اختلاف_ساعت'] > 24).sum())),
        ('🚀 سریع‌ترین تحویل (ساعت)', round(df['اختلاف_ساعت'].min(), 2)),
        ('🐢 کندترین تحویل (ساعت)', round(df['اختلاف_ساعت'].max(), 2)),
        ('📊 چارک اول Q1 (ساعت)', round(kpi.Q1_global, 2)),
        ('📊 چارک سوم Q3 (ساعت)', round(kpi.Q3_global, 2)),
        ('📊 IQR (ساعت)', round(kpi.IQR_global, 2)),
        ('⚠️ آستانه پرت (ساعت)', round(kpi.outlier_threshold, 2)),
        ('⚠️ تعداد مرسولات پرت', len(kpi.df_outliers)),
    ]

    for i, (label, value) in enumerate(kpi_main):
        r = 3 + i
        cell_l = ws.cell(row=r, column=1, value=label)
        cell_l.font = Font(name='B Nazanin', bold=True, size=11)
        cell_l.fill = PatternFill(
            start_color=EXCEL_COLORS['light_blue'],
            end_color=EXCEL_COLORS['light_blue'],
            fill_type='solid',
        )
        cell_l.border = THIN_BORDER
        cell_l.alignment = Alignment(horizontal='right', vertical='center')

        cell_v = ws.cell(row=r, column=2, value=value)
        cell_v.font = Font(name='B Nazanin', bold=True, size=12, color='2F5496')
        cell_v.border = THIN_BORDER
        cell_v.alignment = Alignment(horizontal='center', vertical='center')

    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 22

    # ══════════════════════════════════════════
    # بخش ۲: جداول داده‌ای کوچک (منبع نمودارها)
    # ══════════════════════════════════════════

    # ── جدول ۱: وضعیت تحویل ──
    sc = df['وضعیت_تحویل'].value_counts()
    dr = 19
    ws.cell(row=dr, column=1, value='وضعیت')
    ws.cell(row=dr, column=2, value='تعداد')
    for i, (s, c) in enumerate(sc.items()):
        ws.cell(row=dr + 1 + i, column=1, value=s)
        ws.cell(row=dr + 1 + i, column=2, value=c)
    _style_mini_table(ws, dr, len(sc))

    # ── جدول ۲: بازه‌های تحویل ──
    trd = kpi.time_range_dist
    dr2 = dr + len(sc) + 2
    ws.cell(row=dr2, column=1, value='بازه زمانی')
    ws.cell(row=dr2, column=2, value='تعداد')
    for i, (s, c) in enumerate(trd.items()):
        ws.cell(row=dr2 + 1 + i, column=1, value=s)
        ws.cell(row=dr2 + 1 + i, column=2, value=int(c))
    _style_mini_table(ws, dr2, len(trd))

    # ── جدول ۳: شیفت ورود ──
    ss = kpi.shift_stats
    dr3 = dr2 + len(trd) + 2
    ws.cell(row=dr3, column=1, value='شیفت ورود')
    ws.cell(row=dr3, column=2, value='تعداد')
    for i, (s, row_d) in enumerate(ss.iterrows()):
        ws.cell(row=dr3 + 1 + i, column=1, value=s)
        ws.cell(row=dr3 + 1 + i, column=2, value=int(row_d['تعداد']))
    _style_mini_table(ws, dr3, len(ss))

    # ══════════════════════════════════════════
    # بخش ۳: نمودارهای دایره‌ای
    # ══════════════════════════════════════════
    #
    # چیدمان نهایی (بدون همپوشانی):
    #
    #   ┌──────────────┐  ┌──────────────┐
    #   │  وضعیت تحویل │  │  شیفت ورود   │
    #   │    D3         │  │    N3        │  ← نصف برگشت (بود P3)
    #   └──────────────┘  └──────────────┘
    #   ┌──────────────────┐
    #   │  بازه‌های زمانی   │
    #   │    D28            │  ← ۳ ردیف پایین‌تر (بود D25)
    #   └──────────────────┘

    # ── نمودار ۱: وضعیت تحویل  →  D3 ──
    pie1 = PieChart()
    pie1.title = "نسبت تحویل به‌موقع / تأخیردار"
    pie1.style = 10
    pie1.add_data(
        Reference(ws, min_col=2, min_row=dr, max_row=dr + len(sc)),
        titles_from_data=True,
    )
    pie1.set_categories(
        Reference(ws, min_col=1, min_row=dr + 1, max_row=dr + len(sc))
    )
    pie1.width = 16
    pie1.height = 12
    colorize_pie(pie1, PIE_STATUS_COLORS, len(sc))
    add_data_labels(pie1, show_val=True, show_percent=True, separator="\n")
    _push_pie_plot_down(pie1, offset_y=0.10)
    ws.add_chart(pie1, "D3")

    # ── نمودار ۲: بازه‌های زمانی  →  D28  (بود D25، ۳ ردیف پایین‌تر) ──
    pie2 = PieChart()
    pie2.title = "توزیع بازه‌های زمانی تحویل"
    pie2.style = 10
    pie2.add_data(
        Reference(ws, min_col=2, min_row=dr2, max_row=dr2 + len(trd)),
        titles_from_data=True,
    )
    pie2.set_categories(
        Reference(ws, min_col=1, min_row=dr2 + 1, max_row=dr2 + len(trd))
    )
    pie2.width = 18
    pie2.height = 13
    colorize_pie(pie2, PIE_TIMERANGE_COLORS, len(trd))
    add_data_labels(pie2, show_percent=True, separator="\n")
    _push_pie_plot_down(pie2, offset_y=0.10)
    ws.add_chart(pie2, "D28")   # ← بود D25 → ۳ ردیف پایین‌تر

    # ── نمودار ۳: شیفت ورود  →  N3  (بود P3، نصف برگشت به راست) ──
    # مسیر جابجایی:  L3 (اصلی) → P3 (4 ستون چپ) → N3 (2 ستون برگشت)
    pie3 = PieChart()
    pie3.title = "توزیع شیفت ورود مرسولات"
    pie3.style = 10
    pie3.add_data(
        Reference(ws, min_col=2, min_row=dr3, max_row=dr3 + len(ss)),
        titles_from_data=True,
    )
    pie3.set_categories(
        Reference(ws, min_col=1, min_row=dr3 + 1, max_row=dr3 + len(ss))
    )
    pie3.width = 16
    pie3.height = 12
    add_data_labels(pie3, show_val=True, show_percent=True, separator="\n")
    _push_pie_plot_down(pie3, offset_y=0.10)
    ws.add_chart(pie3, "N3")    # ← بود P3 → نصف برگشت به راست

    return ws
