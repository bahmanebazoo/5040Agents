"""
شیت ۶: تحلیل روز هفته + شیفت
"""
from openpyxl.chart import BarChart, RadarChart, Reference
from openpyxl.chart.label import DataLabelList

from config.constants import WEEKDAY_NAMES
from config.colors import CHART_SPECIFIC
from styles.excel_styles import write_table, auto_width
from charts.chart_helpers import (
    ChartPlacer, colorize_category_bars, apply_single_color, add_data_labels,
)


def build(wb, kpi):
    """ساخت شیت تحلیل روز هفته"""
    ws = wb.create_sheet('تحلیل روز هفته')
    ws.sheet_view.rightToLeft = True

    # ─── جدول روز هفته ───
    headers_wd = ['روز هفته', 'تعداد', 'میانگین (ساعت)', 'درصد به‌موقع']
    rows_wd = []
    for wd in range(7):
        if wd in kpi.weekday_stats.index:
            rd = kpi.weekday_stats.loc[wd]
            rows_wd.append([
                WEEKDAY_NAMES[wd], int(rd['تعداد']),
                rd['میانگین_ساعت'], rd['درصد_بموقع'],
            ])

    lr = write_table(ws, headers_wd, rows_wd, 1, '📆 تحلیل عملکرد به تفکیک روز هفته')
    auto_width(ws, len(headers_wd))

    placer = ChartPlacer(start_row=lr + 1)

    # نمودار ۱: تعداد روز هفته (چندرنگ - دسته‌بندی مفهومی)
    ch1 = BarChart()
    ch1.type = "col"
    ch1.title = "تعداد مرسولات به تفکیک روز هفته"
    ch1.style = 10
    ch1.width = 25
    ch1.height = 15
    ch1.add_data(
        Reference(ws, min_col=2, min_row=2, max_row=lr - 1),
        titles_from_data=True,
    )
    ch1.set_categories(Reference(ws, min_col=1, min_row=3, max_row=lr - 1))
    colorize_category_bars(ch1, len(rows_wd))
    add_data_labels(ch1, show_val=True)
    ws.add_chart(ch1, placer.next_anchor())

    # نمودار ۲: میانگین ساعت روز هفته (چندرنگ)
    ch2 = BarChart()
    ch2.type = "col"
    ch2.title = "میانگین زمان تحویل به تفکیک روز هفته"
    ch2.y_axis.title = "ساعت"
    ch2.style = 10
    ch2.width = 25
    ch2.height = 15
    ch2.add_data(
        Reference(ws, min_col=3, min_row=2, max_row=lr - 1),
        titles_from_data=True,
    )
    ch2.set_categories(Reference(ws, min_col=1, min_row=3, max_row=lr - 1))
    colorize_category_bars(ch2, len(rows_wd))
    add_data_labels(ch2, show_val=True)
    ws.add_chart(ch2, placer.next_anchor())

    # نمودار ۳: Radar
    radar = RadarChart()
    radar.type = "filled"
    radar.title = "نمای رادار عملکرد روزانه"
    radar.style = 10
    radar.width = 25
    radar.height = 15
    radar.add_data(
        Reference(ws, min_col=4, min_row=2, max_row=lr - 1),
        titles_from_data=True,
    )
    radar.set_categories(Reference(ws, min_col=1, min_row=3, max_row=lr - 1))
    ws.add_chart(radar, placer.next_anchor())

    # ─── جدول شیفت ───
    sr_shift = placer.current_row + 2
    headers_sh = ['شیفت', 'تعداد', 'میانگین (ساعت)', 'درصد به‌موقع']
    rows_sh = []
    for sh, rd in kpi.shift_stats.iterrows():
        rows_sh.append([sh, int(rd['تعداد']), rd['میانگین_ساعت'], rd['درصد_بموقع']])

    lr_sh = write_table(
        ws, headers_sh, rows_sh, sr_shift,
        '🕐 تحلیل عملکرد به تفکیک شیفت ورود',
    )

    return ws
