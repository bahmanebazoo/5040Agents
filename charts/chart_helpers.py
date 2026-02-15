"""
charts/chart_helpers.py  — نسخه اصلاح‌شده
"""
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────
# ثوابت تبدیل
# ──────────────────────────────────────────────
# ارتفاع پیش‌فرض هر ردیف اکسل: 15 points = 20 pixels = 0.529 cm
DEFAULT_ROW_HEIGHT_CM = 0.529


def chart_height_to_rows(chart_height_cm: float, gap_rows: int = 2) -> int:
    """
    تبدیل ارتفاع نمودار (سانتیمتر) به تعداد ردیف.
    chart.height در openpyxl بر حسب سانتیمتر است.
    """
    return int(chart_height_cm / DEFAULT_ROW_HEIGHT_CM) + 1


# ──────────────────────────────────────────────
# ChartPlacer — جایگذاری بر اساس سلول
# ──────────────────────────────────────────────
class ChartPlacer:
    """
    جایگذاری نمودارها با فاصله‌گذاری بر اساس واحد سلول.

    استراتژی:
        ۱. ارتفاع واقعی نمودار از chart.height (cm) محاسبه می‌شود
        ۲. فاصله بین نمودارها: gap_rows سلول خالی
        ۳. فرمول: ردیف_بعدی = ردیف_فعلی + ارتفاع_واقعی + gap_rows

    پارامترها:
        start_row    : اولین ردیف مجاز (1-based)
        gap_rows     : تعداد ردیف‌های خالی بین دو نمودار (پیش‌فرض: ۳)
        anchor_col   : ستون شروع نمودار (1-based). 1 = A
    """

    def __init__(
        self,
        start_row: int = 1,
        gap_rows: int = 3,
        anchor_col: int = 1,
        # ── پارامترهای قدیمی (نادیده گرفته می‌شوند — سازگاری عقب‌گرد) ──
        right_margin_px: int = 24,
        gap_px: int = 96,
        default_chart_height_rows: int = 18,
    ):
        self._anchor_col = anchor_col
        self._gap_rows = gap_rows
        self._current_row = start_row
        self._is_first = True

    # ── Properties ──

    @property
    def current_row(self) -> int:
        return self._current_row

    @current_row.setter
    def current_row(self, value: int):
        self._current_row = value

    # ── روش اصلی ──

    def place_chart(self, ws, chart, extra_rows: int = 0):
        """
        جایگذاری نمودار در شیت.

        ارتفاع واقعی نمودار از chart.height (cm) محاسبه می‌شود.
        extra_rows: ردیف‌های اضافی (برای نمودارهایی با Legend بزرگ و...)

        مراحل:
            1. محاسبه ردیف هدف
            2. ws.add_chart(chart, cell_ref)
            3. به‌روزرسانی current_row
        """
        # ردیف هدف
        if self._is_first:
            target_row = self._current_row
        else:
            target_row = self._current_row + self._gap_rows

        # ساخت cell reference
        col_letter = get_column_letter(self._anchor_col)
        cell_ref = f"{col_letter}{target_row}"

        # اضافه کردن نمودار
        ws.add_chart(chart, cell_ref)

        # محاسبه ارتفاع واقعی از chart.height (cm)
        actual_rows = chart_height_to_rows(chart.height) + extra_rows

        # به‌روزرسانی وضعیت
        self._current_row = target_row + actual_rows
        self._is_first = False

    def next_anchor(self, custom_height: int = None) -> str:
        """
        سازگاری عقب‌گرد — فقط cell string برمی‌گرداند.
        ⚠️ این متد ارتفاع واقعی نمودار رو نمی‌دونه.
           اگه ازش استفاده می‌کنید، custom_height رو درست بدید.
        """
        if self._is_first:
            target_row = self._current_row
        else:
            target_row = self._current_row + self._gap_rows

        col_letter = get_column_letter(self._anchor_col)

        chart_h = custom_height or 30  # فرض محافظه‌کارانه
        self._current_row = target_row + chart_h
        self._is_first = False

        return f"{col_letter}{target_row}"


# ──────────────────────────────────────────────
# توابع کمکی نمودار
# ──────────────────────────────────────────────

def apply_single_color(chart, color_hex: str):
    """اعمال رنگ واحد به تمام سری‌های نمودار"""
    for s in chart.series:
        s.graphicalProperties.solidFill = color_hex


def colorize_pie(chart, colors: list, count: int):
    """رنگ‌آمیزی برش‌های نمودار دایره‌ای"""
    if not chart.series:
        return
    for i in range(min(count, len(colors))):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = colors[i]
        chart.series[0].data_points.append(pt)


def colorize_category_bars(chart, num_categories: int):
    """رنگ‌آمیزی هر ستون نمودار میله‌ای با رنگ متفاوت"""
    CATEGORY_PALETTE = [
        "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5",
        "70AD47", "264478", "9B57A0", "636363", "EB7E30",
        "2F5496", "BF9000", "43682B", "C55A11", "7030A0",
        "D34817", "00B0F0", "00B050", "FF0000", "FFC000",
        "8064A2", "F79646", "4BACC6", "C0504D", "9BBB59",
        "8DB4E2", "F4B183", "C5E0B4", "FFD966", "D6B4E0",
        "A9D08E", "BDD7EE", "F8CBAD", "E2EFDA", "FFE699",
    ]
    if not chart.series:
        return
    series = chart.series[0]
    for i in range(num_categories):
        pt = DataPoint(idx=i)
        color = CATEGORY_PALETTE[i % len(CATEGORY_PALETTE)]
        pt.graphicalProperties.solidFill = color
        series.data_points.append(pt)


def add_data_labels(
    chart,
    show_val: bool = False,
    show_percent: bool = False,
    show_cat: bool = False,
    separator: str = "\n",
    font_size: int = 900,
):
    """افزودن برچسب‌های داده به نمودار"""
    if not chart.series:
        return
    dl = DataLabelList()
    dl.showVal = show_val
    dl.showPercent = show_percent
    dl.showCatName = show_cat
    dl.showSerName = False
    dl.separator = separator
    chart.series[0].dLbls = dl
